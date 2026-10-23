from pathlib import Path
import pytest
import openpyxl
from PIL import Image, ImageDraw

from docfast import convert, report
from docfast.util import category


@pytest.fixture
def sample_xlsx(tmp_path) -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Relatório"
    ws.append(["ID", "Item", "Preço", "Quantidade"])
    ws.append(["1", "Cabo USB-C", "29.90", "15"])
    ws.append(["2", "Adaptador HDMI", "59.90", "8"])
    p = tmp_path / "Vendas.xlsx"
    wb.save(p)
    return p


@pytest.fixture
def sample_csv(tmp_path) -> Path:
    p = tmp_path / "Dados.csv"
    p.write_text("Codigo,Nome,Setor\n10,Carlos,TI\n20,Ana,Operações\n", encoding="utf-8")
    return p


@pytest.fixture
def two_images(tmp_path) -> list[Path]:
    paths = []
    for i in range(2):
        im = Image.new("RGB", (100, 100), (50 * i, 100, 150))
        p = tmp_path / f"img_{i}.png"
        im.save(p)
        paths.append(p)
    return paths


def test_sheet_category_and_targets(sample_xlsx, sample_csv):
    assert category(sample_xlsx) == "sheet"
    assert category(sample_csv) == "sheet"

    t_xlsx = convert.targets_for("sheet", ".xlsx")
    assert "PDF" in t_xlsx and "CSV" in t_xlsx and "JSON" in t_xlsx and "HTML" in t_xlsx

    t_csv = convert.targets_for("sheet", ".csv")
    assert "XLSX" in t_csv and "PDF" in t_csv and "JSON" in t_csv and "HTML" in t_csv


def test_xlsx_conversions(sample_xlsx):
    # XLSX -> PDF
    outs_pdf = convert.convert(sample_xlsx, "sheet", "PDF")
    assert len(outs_pdf) == 1 and outs_pdf[0].exists() and outs_pdf[0].suffix == ".pdf"
    assert outs_pdf[0].stat().st_size > 0

    # XLSX -> CSV
    outs_csv = convert.convert(sample_xlsx, "sheet", "CSV")
    assert len(outs_csv) == 1 and outs_csv[0].exists() and outs_csv[0].suffix == ".csv"
    assert "Cabo USB-C" in outs_csv[0].read_text(encoding="utf-8-sig")

    # XLSX -> JSON
    outs_json = convert.convert(sample_xlsx, "sheet", "JSON")
    assert len(outs_json) == 1 and outs_json[0].exists() and outs_json[0].suffix == ".json"
    assert "Cabo USB-C" in outs_json[0].read_text(encoding="utf-8")

    # XLSX -> HTML
    outs_html = convert.convert(sample_xlsx, "sheet", "HTML")
    assert len(outs_html) == 1 and outs_html[0].exists() and outs_html[0].suffix == ".html"
    assert "<table" in outs_html[0].read_text(encoding="utf-8")


def test_csv_conversions(sample_csv):
    # CSV -> XLSX
    outs_xlsx = convert.convert(sample_csv, "sheet", "XLSX")
    assert len(outs_xlsx) == 1 and outs_xlsx[0].exists() and outs_xlsx[0].suffix == ".xlsx"
    wb = openpyxl.load_workbook(outs_xlsx[0])
    assert wb.active.cell(row=2, column=2).value == "Carlos"

    # CSV -> PDF
    outs_pdf = convert.convert(sample_csv, "sheet", "PDF")
    assert len(outs_pdf) == 1 and outs_pdf[0].exists() and outs_pdf[0].suffix == ".pdf"


def test_multiple_images_targets_and_conversion(two_images):
    # Single file targets: should have "PDF"
    single_t = convert.targets_for("image", ".png", num_files=1)
    assert "PDF" in single_t
    assert "PDF (1 doc)" not in single_t

    # Multiple files targets: should have both "PDF (1 doc)" and "PDF (sep)"
    multi_t = convert.targets_for("image", ".png", num_files=len(two_images))
    assert "PDF (1 doc)" in multi_t
    assert "PDF (sep)" in multi_t

    # 1. Combined into 1 PDF
    combined = convert.images_to_single_pdf(two_images)
    assert len(combined) == 1 and combined[0].exists()
    r1 = report.summarize("convert:PDF (1 doc)", two_images, combined)
    assert "2 imagens → 1 PDF" in r1.title

    # 2. Separate PDFs
    seps = [o for p in two_images for o in convert.convert(p, "image", "PDF (sep)")]
    assert len(seps) == 2 and all(p.exists() for p in seps)
    r2 = report.summarize("convert:PDF (sep)", two_images, seps)
    assert "2 imagens → 2 PDFs individuais" in r2.title
